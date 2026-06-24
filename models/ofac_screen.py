"""Automated OFAC SDN screen and Word (.docx) report builder."""

import base64
import csv
import io
import logging
import os
import subprocess
import tempfile
import urllib.request
from datetime import datetime

from odoo import _, api, fields, models
from odoo.exceptions import UserError

from .membership_application import name_similarity

_logger = logging.getLogger(__name__)

OFAC_URL = "https://www.treasury.gov/ofac/downloads/sdn.csv"
OFAC_CACHE_HOURS = 24
SDN_COLS = ["ent_num", "sdn_name", "sdn_type", "program", "title", "call_sign",
            "vess_type", "tonnage", "grt", "vess_flag", "vess_owner", "remarks"]


class OfacScreen(models.AbstractModel):
    _name = "elks.ofac.screen"
    _description = "OFAC SDN automated screen"

    SIMILARITY_THRESHOLD = 0.85
    CACHE_ATTACHMENT_NAME = "elks_ofac_sdn.csv"

    @api.model
    def _download_sdn_csv(self):
        """Try urllib first; fall back to curl (uses macOS/Linux system trust)."""
        try:
            req = urllib.request.Request(
                OFAC_URL,
                headers={"User-Agent": "Odoo-Elks-Investigation-Workspace/1.0"},
            )
            with urllib.request.urlopen(req, timeout=30) as resp:
                data = resp.read()
            if len(data) < 100_000:
                raise RuntimeError(f"file too small ({len(data)} bytes)")
            return data
        except Exception as exc:
            _logger.warning("OFAC urllib failed (%s); falling back to curl", exc)
        with tempfile.NamedTemporaryFile(delete=False) as tmp:
            tmp_path = tmp.name
        try:
            subprocess.run(
                ["curl", "-sSL", "--max-time", "60",
                 "-A", "Odoo-Elks-Investigation-Workspace/1.0",
                 "-o", tmp_path, OFAC_URL],
                check=True,
            )
            with open(tmp_path, "rb") as f:
                data = f.read()
            if len(data) < 100_000:
                raise RuntimeError(f"curl returned file of size {len(data)}")
            return data
        finally:
            try:
                os.unlink(tmp_path)
            except OSError:
                pass

    @api.model
    def _get_sdn_csv(self):
        Attachment = self.env["ir.attachment"].sudo()
        existing = Attachment.search([
            ("name", "=", self.CACHE_ATTACHMENT_NAME),
            ("res_model", "=", False),
        ], limit=1)
        if existing:
            age_h = (datetime.now() - fields.Datetime.from_string(existing.write_date)).total_seconds() / 3600
            if age_h < OFAC_CACHE_HOURS:
                _logger.info("Using cached OFAC SDN (age %.1fh)", age_h)
                return base64.b64decode(existing.datas)
        _logger.info("Refreshing OFAC SDN cache")
        data = self._download_sdn_csv()
        if existing:
            existing.write({"datas": base64.b64encode(data)})
        else:
            Attachment.create({
                "name": self.CACHE_ATTACHMENT_NAME,
                "type": "binary",
                "datas": base64.b64encode(data),
                "res_model": False,
                "res_id": False,
            })
        return data

    @api.model
    def _individuals(self):
        raw = self._get_sdn_csv()
        text = raw.decode("latin-1", errors="replace")
        out = []
        for row in csv.reader(io.StringIO(text)):
            if len(row) < 3:
                continue
            if (row[2] or "").strip().lower() != "individual":
                continue
            padded = row + [""] * (len(SDN_COLS) - len(row))
            out.append(dict(zip(SDN_COLS, padded)))
        return out

    @api.model
    def screen_application(self, application):
        """Run the OFAC SDN screen for an elks.membership.application.

        Returns a dict with keys:
          hits     — list of match dicts (empty when none)
          summary  — human-readable summary line
          searched — count of individuals scanned (0 on download failure)
          failed   — True if download/parse failed
          error    — error message when failed=True
        """
        try:
            individuals = self._individuals()
        except Exception as exc:
            # T4 — graceful failure: caller creates a portal_unavailable
            # check row instead of bubbling a UserError that interrupts the
            # workflow.
            _logger.warning("OFAC screen failed for %s: %s", application.name, exc)
            return {
                "hits": [],
                "summary": f"OFAC CSV download failed: {exc}",
                "searched": 0,
                "failed": True,
                "error": str(exc),
            }
        names = [application.applicant_display_name or ""]
        if application.applicant_first_name and application.applicant_last_name:
            names.append(f"{application.applicant_first_name} {application.applicant_last_name}")
        if application.applicant_maiden_name and application.applicant_first_name:
            names.append(f"{application.applicant_first_name} {application.applicant_maiden_name}")
        names += [
            x.strip() for x in (application.applicant_aliases or "").split(",") if x.strip()
        ]
        names = [n for n in names if n]
        hits = []
        for rec in individuals:
            sdn = rec.get("sdn_name", "")
            if not sdn:
                continue
            best = max(name_similarity(sdn, n) for n in names) if names else 0
            if best >= self.SIMILARITY_THRESHOLD:
                hits.append({
                    "sdn_name": sdn,
                    "ent_num": rec.get("ent_num", ""),
                    "program": rec.get("program", ""),
                    "title": rec.get("title", ""),
                    "remarks": rec.get("remarks", ""),
                    "score": best,
                })
        hits.sort(key=lambda h: -h["score"])
        summary = (f"Searched {len(individuals):,} OFAC individuals with similarity "
                   f">= {self.SIMILARITY_THRESHOLD:.2f}; {len(hits)} match(es).")
        return {
            "hits": hits,
            "summary": summary,
            "searched": len(individuals),
            "failed": False,
            "error": None,
        }


# ---------------------------------------------------------------------------
# Word (.docx) report builder
# ---------------------------------------------------------------------------
class DocxReportBuilder(models.AbstractModel):
    _name = "elks.investigation.docx.builder"
    _description = "DOCX report builder for Elks membership investigations"

    @api.model
    def build_attachment_action(self, application):
        try:
            from docx import Document
            from docx.shared import Pt, RGBColor, Inches
            from docx.enum.text import WD_ALIGN_PARAGRAPH
            from docx.oxml.ns import qn
            from docx.oxml import OxmlElement
        except ImportError:
            raise UserError(_(
                "python-docx is not installed in Odoo's Python environment. "
                "Install it with: pip install python-docx"
            ))

        NAVY = RGBColor(0x1F, 0x38, 0x64)
        GREY = RGBColor(0x59, 0x59, 0x59)

        def shade(cell, hexfill):
            tc_pr = cell._tc.get_or_add_tcPr()
            shd = OxmlElement("w:shd")
            shd.set(qn("w:val"), "clear")
            shd.set(qn("w:color"), "auto")
            shd.set(qn("w:fill"), hexfill)
            tc_pr.append(shd)

        doc = Document()
        for section in doc.sections:
            section.page_width = Inches(8.5)
            section.page_height = Inches(11)
            section.top_margin = Inches(1.0)
            section.bottom_margin = Inches(0.85)
            section.left_margin = Inches(0.75)
            section.right_margin = Inches(0.75)
        doc.styles["Normal"].font.name = "Calibri"
        doc.styles["Normal"].font.size = Pt(11)

        def H(text, size=18, color=NAVY, before=10, after=4):
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(before)
            p.paragraph_format.space_after = Pt(after)
            r = p.add_run(text)
            r.bold = True
            r.font.size = Pt(size)
            r.font.color.rgb = color
            return p

        def P(text, italic=False, color=None, size=11, align=None, after=4):
            p = doc.add_paragraph()
            p.paragraph_format.space_after = Pt(after)
            if align is not None:
                p.alignment = align
            r = p.add_run(text or "")
            r.italic = italic
            r.font.size = Pt(size)
            if color:
                r.font.color.rgb = color
            return p

        def KV(rows, c1=2.2):
            t = doc.add_table(rows=len(rows), cols=2)
            t.style = "Table Grid"
            t.autofit = False
            for i, (k, v) in enumerate(rows):
                c0, c1c = t.rows[i].cells[0], t.rows[i].cells[1]
                c0.width = Inches(c1)
                c1c.width = Inches(7.0 - c1)
                shade(c0, "F2F2F2")
                rk = c0.paragraphs[0].add_run(k)
                rk.bold = True
                rk.font.size = Pt(10)
                rv = c1c.paragraphs[0].add_run(str(v) if v else "")
                rv.font.size = Pt(10)
            return t

        def TBL(headers, rows, widths):
            t = doc.add_table(rows=1 + len(rows), cols=len(headers))
            t.style = "Table Grid"
            t.autofit = False
            for i, h in enumerate(headers):
                c = t.rows[0].cells[i]
                c.width = Inches(widths[i])
                shade(c, "1F3864")
                r = c.paragraphs[0].add_run(h)
                r.bold = True
                r.font.size = Pt(10)
                r.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
            for ri, row in enumerate(rows, start=1):
                for ci, val in enumerate(row):
                    c = t.rows[ri].cells[ci]
                    c.width = Inches(widths[ci])
                    if ri % 2 == 0:
                        shade(c, "F7F7F7")
                    r = c.paragraphs[0].add_run(str(val) if val else "")
                    r.font.size = Pt(10)
            return t

        a = application

        P("BENEVOLENT AND PROTECTIVE ORDER OF ELKS — LEWISTON LODGE #896",
          color=NAVY, size=11, align=WD_ALIGN_PARAGRAPH.CENTER, after=2)
        P("Investigating Committee — Background Investigation Findings Report",
          italic=True, color=GREY, size=11, align=WD_ALIGN_PARAGRAPH.CENTER, after=12)

        ct = doc.add_table(rows=1, cols=1)
        ct.style = "Table Grid"
        cell = ct.rows[0].cells[0]
        shade(cell, "FFF2CC")
        cp = cell.paragraphs[0]
        cp.add_run("CONFIDENTIAL — INVESTIGATING COMMITTEE EYES ONLY. ").bold = True
        cp.add_run("This report consolidates findings from official government records, "
                   "compiled manually by the committee. For internal "
                   "membership-investigation use only.")
        doc.add_paragraph()

        H("1.  Applicant Summary", size=16)
        KV([
            ("Application #", a.name),
            ("Application type", dict(a._fields["application_type"].selection).get(a.application_type, "")),
            ("Stage", dict(a._fields["stage"].selection).get(a.stage, "")),
            ("Date proposed", str(a.date_proposed or "")),
            ("Investigator", a.investigator_id.name or ""),
            ("Investigation period",
             f"{a.date_investigation_assigned or ''} — {a.date_investigation_complete or '(in progress)'}"),
            ("Applicant", a.applicant_display_name),
            ("Maiden / aliases", ", ".join(filter(None, [a.applicant_maiden_name, a.applicant_aliases])) or "—"),
            ("DOB", str(a.applicant_date_of_birth or "—")),
            ("Current address",
             ", ".join(filter(None, [a.applicant_street, a.applicant_city,
                                      a.applicant_state_id.code if a.applicant_state_id else None,
                                      a.applicant_zip]))),
            ("Phone / Email", f"{a.applicant_phone or '—'}  /  {a.applicant_email or '—'}"),
        ])

        H("2.  Residential History (10-year)", size=14)
        if a.investigation_address_history_ids:
            TBL(["Street", "City", "State", "From", "To"],
                [(h.street, h.city, h.state_code or '', h.date_from, h.date_to)
                 for h in a.investigation_address_history_ids],
                [2.6, 1.6, 0.7, 1.0, 1.1])
        else:
            P("(none recorded — applicant has only the current address on file)",
              italic=True, color=GREY)

        H("3.  Employment History (10-year)", size=14)
        if a.investigation_employment_history_ids:
            TBL(["Employer", "City", "State", "Role", "From", "To"],
                [(e.employer, e.city, e.state_code or '', e.role or '', e.date_from, e.date_to)
                 for e in a.investigation_employment_history_ids],
                [1.8, 1.3, 0.6, 1.5, 0.9, 0.9])
        else:
            P("(none recorded)", italic=True, color=GREY)

        H("4.  Reference Checks", size=14)
        if a.investigation_reference_ids:
            TBL(["Reference", "Relationship", "Date", "Recommendation", "Notes"],
                [(r.name, r.relationship or '', str(r.contacted_on or ''),
                  dict(r._fields["recommendation"].selection).get(r.recommendation, '')
                    if r.recommendation else '',
                  r.notes or '')
                 for r in a.investigation_reference_ids],
                [1.5, 1.3, 0.9, 1.5, 1.8])
        else:
            P("(no references recorded)", italic=True, color=GREY)

        doc.add_page_break()
        H("5.  Jurisdiction-by-Jurisdiction Findings", size=16)
        by_juris = {}
        for chk in a.investigation_check_ids:
            by_juris.setdefault(chk.jurisdiction_code, []).append(chk)
        for juris in (["NATIONAL"] + a._investigation_state_codes()):
            if juris not in by_juris:
                continue
            H(juris, size=13)
            for chk in by_juris[juris]:
                P(chk.source_label or "", italic=False, after=2)
                KV([
                    ("Source", chk.source_code),
                    ("URL", chk.url or ""),
                    ("Date checked", str(chk.date_checked or "")),
                    ("Operator", chk.operator_id.name or ""),
                    ("Result", dict(chk._fields["result_summary"].selection).get(chk.result_summary, "")),
                    ("Notes", chk.operator_notes or "—"),
                ], c1=1.6)
                if chk.hit_ids:
                    P("Records captured:", italic=True, color=GREY, after=2)
                    TBL(["Name on record", "DOB", "Case #", "Charges", "Disposition", "Notes"],
                        [(h.name_on_record, h.dob_on_record, h.case_number, h.charges or "",
                          h.disposition or "", h.notes or "")
                         for h in chk.hit_ids],
                        [1.5, 0.9, 1.0, 1.5, 1.1, 1.0])
                doc.add_paragraph()

        doc.add_page_break()
        H("6.  Cross-Check Flags", size=16)
        if not a.investigation_flag_ids:
            P("No anomalies detected. Identity, dates, addresses, and coverage are internally consistent.")
        else:
            for sev in ("high", "medium", "low", "info"):
                fs = a.investigation_flag_ids.filtered(lambda f: f.severity == sev)
                if not fs:
                    continue
                color = {"high": RGBColor(0xC0, 0x39, 0x2B),
                         "medium": RGBColor(0xD9, 0x82, 0x2A),
                         "low": RGBColor(0x1F, 0x77, 0xB4),
                         "info": GREY}[sev]
                H(f"{sev.upper()} severity", size=12, color=color, before=8, after=2)
                for f in fs:
                    P(f"• ({f.category}) {f.message}", size=10)

        H("7.  Committee Recommendation", size=16)
        if a.investigation_result:
            color = {"favorable": RGBColor(0x2E, 0x86, 0x40),
                     "unfavorable": RGBColor(0xC0, 0x39, 0x2B)}[a.investigation_result]
            rp = doc.add_paragraph()
            r = rp.add_run(
                f"Recommendation:  {'APPROVE' if a.investigation_result == 'favorable' else 'DENY'}"
            )
            r.bold = True
            r.font.size = Pt(14)
            r.font.color.rgb = color
        if a.investigation_result == "unfavorable" and a.investigation_denial_reason:
            H("Reason for denial", size=12, color=GREY, before=10, after=4)
            P(a.investigation_denial_reason, after=8)
        if a.investigation_notes:
            H("Investigator notes", size=12, color=GREY, before=10, after=4)
            P(a.investigation_notes)

        H("8.  Adverse-Action Procedure (if recommendation is DENY)", size=14)
        P("If the membership decision is AGAINST and was based on a third-party "
          "consumer-reporting agency report, the Lodge must follow the federal "
          "FCRA two-step adverse-action procedure: pre-adverse-action notice "
          "with a copy of the report and the FCRA summary-of-rights, allowing "
          "a reasonable time (typically 5 business days) to dispute; then a "
          "final adverse-action notice identifying the reporting agency and "
          "the applicant's right to dispute the report directly.")

        buf = io.BytesIO()
        doc.save(buf)
        data = buf.getvalue()

        fname = f"Investigation_{a.name}_{(a.applicant_display_name or 'applicant').replace(' ', '_')}.docx"
        att = self.env["ir.attachment"].create({
            "name": fname,
            "type": "binary",
            "datas": base64.b64encode(data),
            "res_model": "elks.membership.application",
            "res_id": a.id,
            "mimetype": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        })
        return {
            "type": "ir.actions.act_url",
            "url": f"/web/content/{att.id}?download=true",
            "target": "self",
        }
